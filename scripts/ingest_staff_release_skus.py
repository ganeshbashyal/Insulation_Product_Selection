"""Prepare a local, non-authoritative SKU inventory from the staff release XLSM.

The Melbourne sales tab is the baseline. Other state tabs are retained only as
coverage/variance annotations; prices are compared locally but never written to
the output artifact. Family links are candidates for authoring, not approvals.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts.ingest_product_master import load_cards, match_groups  # noqa: E402

SALES_FIELDS = (
    "Our SKU", "Our Product Name", "Category", "Material Type", "Product Use",
    "R Value / RW / NRC", "Length (mm)", "Width (mm)", "Thickness (mm)",
    "M2 per Unit", "Buy/Sell Unit", "Sell Price (Inc GST)", "MOQ",
)
COMPARE_FIELDS = (
    "Our Product Name", "Category", "Material Type", "Product Use",
    "R Value / RW / NRC", "Length (mm)", "Width (mm)", "Thickness (mm)",
    "M2 per Unit", "Buy/Sell Unit", "MOQ",
)
CURRENT_MIGRATION_STATUSES = {"RENAMED", "ADDED", "NEW"}
FALLBACK_MIGRATION_STATUSES = {"MERGED", "CONSOLIDATED"}


def normalized(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def clean(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def price_status(value: object) -> str:
    if value is None or not clean(value):
        return "blank"
    if isinstance(value, (int, float)):
        return "numeric"
    if clean(value).casefold() == "poa":
        return "poa"
    return "other"


def _headers(row: tuple, sheet: str, required: set[str]) -> dict[str, int]:
    positions = {clean(value): index for index, value in enumerate(row) if clean(value)}
    missing = sorted(required - positions.keys())
    if missing:
        raise ValueError(f"{sheet} is missing required columns: {', '.join(missing)}")
    return positions


def _read_sales_sheet(sheet, location: str) -> dict[str, dict]:
    iterator = sheet.iter_rows(values_only=True)
    next(iterator, None)
    next(iterator, None)
    header = next(iterator, None)
    if header is None:
        raise ValueError(f"{sheet.title} has no sales header row")
    positions = _headers(header, sheet.title, set(SALES_FIELDS))
    rows: dict[str, dict] = {}
    for row_number, values in enumerate(iterator, 4):
        values = tuple(values)
        sku = clean(values[positions["Our SKU"]] if len(values) > positions["Our SKU"] else "")
        if not sku:
            continue
        key = sku.casefold()
        if key in rows:
            raise ValueError(f"{sheet.title} has duplicate SKU {sku!r}")
        item = {field: values[positions[field]] if len(values) > positions[field] else None
                for field in SALES_FIELDS}
        item.update({"sku": sku, "location": location, "sheet_row": row_number})
        rows[key] = item
    return rows


def _read_product_master(sheet) -> dict[str, dict]:
    iterator = sheet.iter_rows(values_only=True)
    next(iterator, None)
    header = next(iterator, None)
    if header is None:
        raise ValueError("Product_Master has no header row")
    required = {"Our SKU", "Our Product Name", "Active?", "Manufacturer Name",
                "Category", "Material Type", "Product Use", "Spec Family Name"}
    positions = _headers(header, sheet.title, required)
    products: dict[str, dict] = {}
    for row_number, values in enumerate(iterator, 3):
        values = tuple(values)
        sku = clean(values[positions["Our SKU"]] if len(values) > positions["Our SKU"] else "")
        if not sku:
            continue
        key = sku.casefold()
        if key in products:
            raise ValueError(f"Product_Master has duplicate SKU {sku!r}")
        get = lambda name: values[positions[name]] if len(values) > positions[name] else None
        products[key] = {
            "sku": sku,
            "product_name": clean(get("Our Product Name")),
            "active": clean(get("Active?")).casefold() == "yes",
            "manufacturer": clean(get("Manufacturer Name")),
            "category": clean(get("Category")),
            "material_type": clean(get("Material Type")),
            "spec_material_type": clean(get("Spec Material Type")) if "Spec Material Type" in positions else "",
            "product_use": clean(get("Product Use")),
            "spec_family_name": clean(get("Spec Family Name")),
            "spec_id": clean(get("Spec ID")) if "Spec ID" in positions else "",
            "source_row": row_number,
        }
    return products


def _read_migration_register(sheet, active_skus: set[str]) -> dict[str, dict]:
    iterator = sheet.iter_rows(values_only=True)
    for _ in range(3):
        next(iterator, None)
    header = next(iterator, None)
    if header is None:
        raise ValueError("SKU_Migration_Register has no header row")
    required = {"Status", "New / Canonical SKU", "Manufacturer", "Family"}
    positions = _headers(header, sheet.title, required)
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row_number, values in enumerate(iterator, 5):
        values = tuple(values)
        get = lambda name: values[positions[name]] if len(values) > positions[name] else None
        sku = clean(get("New / Canonical SKU"))
        if not sku or sku.casefold() not in active_skus:
            continue
        grouped[sku.casefold()].append({
            "status": clean(get("Status")),
            "family_code": clean(get("Family")),
            "manufacturer": clean(get("Manufacturer")),
            "row": row_number,
        })
    result = {}
    for sku, entries in grouped.items():
        preferred = [entry for entry in entries if entry["status"] in CURRENT_MIGRATION_STATUSES]
        if not preferred:
            preferred = [entry for entry in entries if entry["status"] in FALLBACK_MIGRATION_STATUSES]
        if not preferred:
            preferred = entries
        result[sku] = {
            "entries": entries,
            "selected_entries": preferred,
            "family_codes": sorted({entry["family_code"] for entry in preferred if entry["family_code"]}),
            "statuses": sorted({entry["status"] for entry in entries if entry["status"]}),
        }
    return result


def _family_groups(products: list[dict], migrations: dict[str, dict]) -> list[dict]:
    grouped: dict[tuple, list[dict]] = defaultdict(list)
    for product in products:
        migration = migrations.get(product["sku"].casefold(), {})
        family_codes = migration.get("family_codes", [])
        code = " / ".join(family_codes)
        key = (
            normalized(product["manufacturer"]), code, normalized(product["spec_family_name"]),
            normalized(product["product_use"]), normalized(product["category"]),
        )
        grouped[key].append(product)
    groups = []
    for key, items in sorted(grouped.items()):
        manufacturer = items[0]["manufacturer"]
        code = " / ".join(sorted({code for item in items
                                  for code in migrations.get(item["sku"].casefold(), {}).get("family_codes", [])}))
        spec_names = sorted({item["spec_family_name"] for item in items if item["spec_family_name"]})
        uses = sorted({item["product_use"] for item in items if item["product_use"]})
        categories = sorted({item["category"] for item in items if item["category"]})
        group_id = hashlib.sha256(json.dumps(key, ensure_ascii=True).encode()).hexdigest()[:20]
        groups.append({
            "group_id": group_id,
            "manufacturer": manufacturer,
            "family_code": code,
            "spec_family_names": spec_names,
            "product_uses": uses,
            "categories": categories,
            "product_count": len(items),
            "product_names": sorted({item["product_name"] for item in items if item["product_name"]})[:12],
            "sku_codes": sorted({item["sku"] for item in items}, key=str.casefold),
            "items": items,
        })
    return groups


def _deterministic_matches(groups: list[dict]) -> dict[str, dict]:
    import pandas as pd

    rows = []
    for group in groups:
        rows.append({
            "group_id": group["group_id"],
            "manufacturername": group["manufacturer"],
            "specfamilyname": " | ".join([
                group["family_code"], *group["spec_family_names"],
                *group["product_uses"], *group["categories"],
            ]),
            "ourproductname": " | ".join(group["product_names"]),
            "category": " | ".join(group["categories"]),
            "productuse": " | ".join(group["product_uses"]),
            "materialtype": " | ".join(sorted({
                item["material_type"] for item in group["items"] if item["material_type"]
            })),
            "specmaterialtype": " | ".join(sorted({
                item["spec_material_type"] for item in group["items"] if item["spec_material_type"]
            })),
        })
    decisions = match_groups(pd.DataFrame(rows), load_cards())
    return {
        row["group_id"]: decisions[(normalized(row["manufacturername"]), row["specfamilyname"])]
        for row in rows
    }


def apply_family_map_to_inventory(inventory: dict, family_map: dict) -> dict:
    if family_map.get("source_sha256") != inventory.get("source_sha256"):
        raise ValueError("Local model map is bound to a different workbook SHA-256")
    if inventory.get("price_values_included") is not False:
        raise ValueError("Existing SKU inventory is not marked price-redacted")
    proposals = family_map.get("groups")
    if not isinstance(proposals, dict):
        raise ValueError("Local model map groups must be an object")

    requests = {row["group_id"]: row for row in inventory.get("mapping_requests", [])}
    groups = {row["group_id"]: row for row in inventory.get("groups", [])}
    if not proposals or not set(proposals).issubset(requests) or not set(requests).issubset(groups):
        raise ValueError("Local model map must contain valid inventory review groups")
    products_by_group = defaultdict(list)
    for product in inventory.get("products", []):
        products_by_group[product.get("family_mapping_group")].append(product)

    for group_id, proposal in proposals.items():
        if not isinstance(proposal, dict):
            raise ValueError(f"Malformed local model decision for group {group_id}")
        confidence = proposal.get("confidence")
        if isinstance(confidence, bool) or not isinstance(confidence, int) or not 0 <= confidence <= 100:
            raise ValueError(f"Invalid local model confidence for group {group_id}")
        proposed_id = proposal.get("family_id")
        request = requests[group_id]
        candidates = {item["family_id"]: item for item in request.get("candidates", [])}
        if proposed_id is not None:
            if not isinstance(proposed_id, str) or proposed_id not in candidates:
                raise ValueError(f"Local model selected a non-candidate family for group {group_id}")
            if normalized(candidates[proposed_id].get("manufacturer")) != normalized(request.get("manufacturer")):
                raise ValueError(f"Local model crossed manufacturers for group {group_id}")
        reason = proposal.get("reason", "")
        if not isinstance(reason, str):
            raise ValueError(f"Invalid local model reason for group {group_id}")

        group = groups[group_id]
        deterministic_family = group.get("family_id")
        candidate_id = deterministic_family or proposed_id or group.get("candidate_family_id")
        if deterministic_family:
            status = "deterministic_high_unreviewed"
        elif proposed_id:
            status = "local_model_candidate_unreviewed"
        elif candidate_id:
            status = "heuristic_candidate_review_required"
        else:
            status = "unmapped_review_required"
        group.update({
            "candidate_family_id": candidate_id,
            "mapping_status": status,
            "model_confidence": confidence,
            "model_reason": clean(reason),
        })
        for product in products_by_group.get(group_id, []):
            product["candidate_family_id"] = candidate_id
            product["family_mapping_status"] = status
            product["model_confidence"] = confidence

    inventory["family_mapping_model"] = {
        "name": clean(family_map.get("model")) or "local model",
        "source_sha256": inventory["source_sha256"],
        "review_groups_mapped": len(proposals),
        "review_groups_total": len(requests),
        "complete": set(proposals) == set(requests),
    }
    return inventory


def build_inventory(source: Path, family_map_path: Path | None = None) -> dict:
    source = source.resolve()
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True, keep_vba=False)
    try:
        required_sheets = {"Product_Master", "SKU_Migration_Register",
                           "Sales_Melbourne", "Sales_Sydney", "Sales_Brisbane",
                           "Sales_Adelaide", "Sales_Perth"}
        missing = sorted(required_sheets - set(workbook.sheetnames))
        if missing:
            raise ValueError(f"Workbook is missing required sheets: {', '.join(missing)}")
        product_master = _read_product_master(workbook["Product_Master"])
        active_master = {key: row for key, row in product_master.items() if row["active"]}
        locations = {
            name.removeprefix("Sales_"): _read_sales_sheet(workbook[name], name.removeprefix("Sales_"))
            for name in sorted(required_sheets) if name.startswith("Sales_")
        }
        migrations = _read_migration_register(workbook["SKU_Migration_Register"], set(active_master))
    finally:
        workbook.close()

    union = set().union(*(set(rows) for rows in locations.values()))
    if union != set(active_master):
        missing_sales = sorted(set(active_master) - union)
        unknown_sales = sorted(union - set(active_master))
        raise ValueError(
            f"Sales tabs do not match active Product_Master: {len(missing_sales)} master SKUs missing, "
            f"{len(unknown_sales)} sales-only SKUs"
        )
    baseline = locations["Melbourne"]
    products = []
    state_summary = {}
    for location, rows in locations.items():
        state_summary[location] = {
            "sku_count": len(rows),
            "melbourne_skus_not_listed": len(set(baseline) - set(rows)) if location != "Melbourne" else 0,
            "state_only_skus": len(set(rows) - set(baseline)) if location != "Melbourne" else 0,
            "field_differences": Counter(),
            "sell_price_variances": 0,
        }
    for key, master in active_master.items():
        by_location = {location: rows[key] for location, rows in locations.items() if key in rows}
        baseline_row = by_location.get("Melbourne")
        reference = baseline_row or next(iter(by_location.values()))
        differences = {}
        price_statuses = {}
        for location in locations:
            row = by_location.get(location)
            if row is None:
                differences[location] = ["not_listed"]
                continue
            price_statuses[location] = price_status(row["Sell Price (Inc GST)"])
            if baseline_row is None or location == "Melbourne":
                differences[location] = []
                continue
            changed = []
            for field in COMPARE_FIELDS:
                left = baseline_row[field]
                right = row[field]
                if (left is None or not clean(left)) and (right is None or not clean(right)):
                    continue
                if clean(left) != clean(right):
                    changed.append(field)
                    state_summary[location]["field_differences"][field] += 1
            try:
                price_changed = float(baseline_row["Sell Price (Inc GST)"]) != float(row["Sell Price (Inc GST)"])
            except (TypeError, ValueError):
                price_changed = price_status(baseline_row["Sell Price (Inc GST)"]) != price_status(row["Sell Price (Inc GST)"])
            if price_changed:
                changed.append("sell_price_varies")
                state_summary[location]["sell_price_variances"] += 1
            differences[location] = changed
        migration = migrations.get(key, {"entries": [], "selected_entries": [], "family_codes": [], "statuses": []})
        products.append({
            **master,
            "family_codes": migration["family_codes"],
            "migration_statuses": migration["statuses"],
            "migration_code_conflict": len(migration["family_codes"]) > 1,
            "sales_source_rows": {loc: row["sheet_row"] for loc, row in by_location.items()},
            "melbourne_baseline": baseline_row is not None,
            "state_presence": sorted(by_location),
            "state_differences": differences,
            "price_statuses": price_statuses,
            "price_values_included": False,
            "our_sku": clean(reference["sku"]),
            "sales_product_name": clean(reference["Our Product Name"]),
        })
    for summary in state_summary.values():
        summary["field_differences"] = dict(summary["field_differences"])

    groups = _family_groups(products, migrations)
    matches = _deterministic_matches(groups)
    family_map = {}
    model_info = None
    if family_map_path:
        family_map = json.loads(family_map_path.read_text(encoding="utf-8"))
        if family_map.get("source_sha256") != source_hash:
            raise ValueError("Local model map is bound to a different workbook SHA-256")
        model_info = {"name": clean(family_map.get("model")) or "local model",
                      "source_sha256": source_hash}
        family_map = family_map.get("groups", {})

    cards = load_cards()
    family_info = {}
    for card in cards:
        family_info[card["family_id"]] = {
            "manufacturer": clean(card.get("manufacturer")),
            "name": clean(card.get("name")),
        }
    group_results = []
    for group in groups:
        match = matches[group["group_id"]]
        deterministic_family = match.get("family_id") if match.get("tier") == "high" else None
        proposed = family_map.get(group["group_id"], {})
        proposed_id = proposed.get("family_id") if isinstance(proposed, dict) else None
        if proposed_id and proposed_id not in family_info:
            raise ValueError(f"Local model returned unknown family ID for group {group['group_id']}")
        if proposed_id and normalized(family_info[proposed_id]["manufacturer"]) != normalized(group["manufacturer"]):
            raise ValueError(f"Local model crossed manufacturers for group {group['group_id']}")
        candidate_id = deterministic_family or proposed_id or match.get("candidate_family_id")
        if deterministic_family:
            mapping_status = "deterministic_high_unreviewed"
        elif proposed_id:
            mapping_status = "local_model_candidate_unreviewed"
        elif candidate_id:
            mapping_status = "heuristic_candidate_review_required"
        else:
            mapping_status = "unmapped_review_required"
        decision = {
            "group_id": group["group_id"],
            "manufacturer": group["manufacturer"],
            "family_code": group["family_code"],
            "product_count": group["product_count"],
            "sku_codes": group["sku_codes"],
            "family_id": deterministic_family,
            "candidate_family_id": candidate_id,
            "mapping_status": mapping_status,
            "deterministic_tier": match.get("tier"),
            "deterministic_evidence": match.get("evidence", ""),
            "model_confidence": proposed.get("confidence") if isinstance(proposed, dict) else None,
            "model_reason": clean(proposed.get("reason")) if isinstance(proposed, dict) else "",
            "spec_family_names": group["spec_family_names"],
            "product_uses": group["product_uses"],
        }
        group_results.append(decision)
        for item in group["items"]:
            item["family_id"] = deterministic_family
            item["candidate_family_id"] = candidate_id
            item["family_mapping_status"] = mapping_status
            item["family_mapping_group"] = group["group_id"]
            item["deterministic_tier"] = match.get("tier")
            item["model_confidence"] = decision["model_confidence"]

    request_groups = []
    for group, match in zip(groups, [matches[g["group_id"]] for g in groups]):
        if match.get("tier") == "high":
            continue
        candidates = [
            {"family_id": fid, **family_info[fid]}
            for fid in sorted(family_info)
            if normalized(family_info[fid]["manufacturer"]) == normalized(group["manufacturer"])
        ]
        if candidates:
            request_groups.append({
                "group_id": group["group_id"],
                "manufacturer": group["manufacturer"],
                "family_codes": group["family_code"],
                "spec_family_names": group["spec_family_names"],
                "product_uses": group["product_uses"],
                "categories": group["categories"],
                "product_names": group["product_names"][:8],
                "candidates": candidates,
                "deterministic_candidate": match.get("candidate_family_id"),
                "deterministic_tier": match.get("tier"),
                "deterministic_evidence": match.get("evidence", ""),
            })
    return {
        "schema_version": 1,
        "source_workbook": source.name,
        "source_sha256": source_hash,
        "source_kind": "staff price-list SKU identity and regional coverage",
        "baseline_location": "Melbourne",
        "price_values_included": False,
        "price_variance_method": "Compared locally; output retains only whether/status that a value differs.",
        "active_product_master_skus": len(active_master),
        "inactive_product_master_skus_excluded": len(product_master) - len(active_master),
        "sales_location_summary": state_summary,
        "family_mapping_model": model_info,
        "groups": group_results,
        "mapping_requests": request_groups,
        "products": products,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    inputs = parser.add_mutually_exclusive_group(required=True)
    inputs.add_argument("--source", type=Path)
    inputs.add_argument("--existing-inventory", type=Path,
                        help="apply a complete or partial source-hash-bound map to an existing price-redacted inventory")
    parser.add_argument("--output", type=Path, default=ROOT / "data" / "local" / "staff_release_skus.json")
    parser.add_argument("--family-map", type=Path,
                        help="local model mapping JSON bound to this workbook's SHA-256")
    parser.add_argument("--mapping-requests", type=Path,
                        help="write deterministic review groups and same-manufacturer family candidates")
    args = parser.parse_args()
    if args.existing_inventory:
        if not args.family_map:
            parser.error("--existing-inventory requires --family-map")
        inventory = json.loads(args.existing_inventory.read_text(encoding="utf-8"))
        family_map = json.loads(args.family_map.read_text(encoding="utf-8"))
        inventory = apply_family_map_to_inventory(inventory, family_map)
    else:
        inventory = build_inventory(args.source, args.family_map)
    if args.mapping_requests:
        output = {
            "source_workbook": inventory["source_workbook"],
            "source_sha256": inventory["source_sha256"],
            "groups": inventory["mapping_requests"],
        }
        args.mapping_requests.parent.mkdir(parents=True, exist_ok=True)
        args.mapping_requests.write_text(json.dumps(output, indent=2), encoding="utf-8")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(inventory, indent=2), encoding="utf-8")
    statuses = Counter(row["family_mapping_status"] for row in inventory["products"])
    print(json.dumps({
        "output": str(args.output),
        "source_sha256": inventory["source_sha256"],
        "active_skus": len(inventory["products"]),
        "families_with_candidates": len({row["candidate_family_id"] for row in inventory["products"]
                                         if row["candidate_family_id"]}),
        "mapping_status": dict(statuses),
        "state_summary": inventory["sales_location_summary"],
        "mapping_requests": str(args.mapping_requests) if args.mapping_requests else None,
    }, indent=2))


if __name__ == "__main__":
    main()
