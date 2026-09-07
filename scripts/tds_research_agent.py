"""Deep-dive research agent: fetch real manufacturer datasheets and populate
rich product literature for every family.

Pipeline per family (runs locally, resumable):
  1. resolve   - find the TDS/SDS URL: use the confirmed source_url, else
                 search the manufacturer's own site (DuckDuckGo site: query,
                 no API key) for a datasheet PDF on the official domain
  2. fetch     - download the PDF to data/local/tds_cache/ (cached by URL)
  3. extract   - pull text with pypdf
  4. structure - ask the LOCAL Ollama model to extract a strict JSON spec
                 (description, features, technical table, applications,
                 install, fire, sustainability). No data leaves the machine.
  5. write     - store the structured spec to knowledge/<mfg>/research/<slug>.json

After research, run scripts/generate_family_literature.py --use-research to
regenerate the MD/DOCX from the real extracted data instead of thin catalogue
placeholders.

Usage:
    python scripts/tds_research_agent.py --only Thermotec     # one manufacturer
    python scripts/tds_research_agent.py --limit 5            # first 5 families
    python scripts/tds_research_agent.py                      # all, resumable
    python scripts/tds_research_agent.py --status             # progress report
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FutureTimeoutError
from pathlib import Path

import pypdf
import requests

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import llm_client
from audit_datasheet_links import OFFICIAL_DOMAINS, domain_matches, domain_of

# Extraction is pure text-restructuring (temperature 0, no creative reasoning).
# phi4-mini benchmarked fastest locally on this CPU-only box (~68s/family vs
# ~93s for gemma4 and 350s+ for qwen3:8b, which is unexpectedly slow here
# despite think:false). Override with OLLAMA_EXTRACT_MODEL if that changes.
EXTRACT_MODEL = os.getenv("OLLAMA_EXTRACT_MODEL", "phi4-mini:latest")

CACHE_DIR = ROOT / "data" / "local" / "tds_cache"
RESEARCH_DIR_NAME = "research"
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) tds-research/1.0"

BASE_EXTRACT_PROMPT = """You are extracting structured product data from an insulation Technical Data Sheet.

Return ONLY a JSON object (no markdown fences, no commentary) with these keys:
  "description": 1-2 sentence factual product description.
  "features": list of up to 8 short feature strings.
  "applications": list of applications (e.g. "ceiling", "external wall").
  "technical": list of {"property","value","standard"} objects for every general spec found (R-value, density, thermal conductivity, fire indices, temperature range, vapour, pH). Use "" for standard if none stated.
  "fire": short string of fire/AS-NZS 1530.3 results, or "" if none.
  "sustainability": short string of recycled-content/VOC/environmental claims, or "" if none.
  "install": list of up to 12 short installation steps actually described in the text (fixing method, spacing, compression/gap avoidance, vapour barrier orientation, handling), or [] if none.
  "clearances": list of up to 8 short strings describing required clearances/safe distances (downlights, flues, exhaust fans, electrical) if stated, or [] if none.
  "limitations": list of up to 6 short manufacturer-stated limitations or warnings, or [] if none.
{range_fields}
Rules: only report values actually present in the text. Never invent numbers or rows. If a field is absent use "" or []. Keep numbers, units and product codes exactly as written.

TDS TEXT:
"""

# Appended to BASE_EXTRACT_PROMPT only when the deterministic regex table parser
# (parse_variant_table) found no variant table, so the model isn't asked to do
# work the parser already does faster and more reliably.
RANGE_FIELDS_PROMPT = """  "range_headers": the exact column headings of the product's physical characteristics / dimensions / packaging table (e.g. ["Material R-value","Nominal thickness (mm)","Width (mm)","Length (mm)","Batts per pack","m2 per pack","Coverage per pack (m2)","Packs per bale","Product code"]). [] if no such table exists.
  "range": list of row objects, one per size/variant/product code in that table, each shaped {"c0":...,"c1":...} matching range_headers by position. EVERY row must have exactly one key per header, in order, with no gaps — if the source table merges a cell (e.g. the same R-value or thickness spans two width rows), repeat that value in both rows rather than omitting it. Combine a value and its unit into one header/cell (e.g. "R2.5", not separate "R-value" and "m2 K/W" headers). Copy EVERY row present — do not summarize, average or omit any width, thickness or product code variant. [] if none. The range table must be exhaustive: if the source table has 16 rows, return 16 rows.
"""


def slugify(name: str) -> str:
    return re.sub(r"[^a-zA-Z0-9]+", "_", name).strip("_").lower()[:60]


# Deterministic parser for the common AU glasswool batt "Physical Characteristics"
# table layout: a full row states R-value, thickness, width, length, batts/pack,
# m2/pack, coverage/pack and product code; a following row often re-states only
# the width variant at the same R-value/thickness (a merged cell in the source
# PDF), omitting the first two fields. Parsing this with regex is instant and
# never misaligns columns the way an LLM reproduction of a merged-cell table can.
_BATT_TABLE_HEADERS = [
    "Material R-value", "Nominal thickness (mm)", "Width (mm)", "Length (mm)",
    "Batts per pack", "m2 per pack", "Coverage per pack (m2)", "Packs per bale", "Product code",
]
_BATT_FULL_ROW = re.compile(
    r"^\s*(R\d+(?:\.\d+)?)\s+(\d{2,4})\s+(\d{2,4})\s+(\d{3,5})\s+(\d{1,3})\s+([\d.]+)\s+([\d.]+)\s+(\d{1,3})\s+(\d{5,8})\s*$",
    re.M,
)
_BATT_CONT_ROW = re.compile(
    r"^\s*(\d{2,4})\s+(\d{3,5})\s+(\d{1,3})\s+([\d.]+)\s+([\d.]+)\s+(\d{1,3})\s+(\d{5,8})\s*$",
    re.M,
)
_TABLE_SECTION = re.compile(r"physical (?:characteristics|properties)", re.I)


def parse_variant_table(text: str) -> tuple[list[str], list[dict]] | None:
    """Best-effort deterministic extraction of the batt dimensions/packaging
    table. Returns (headers, rows) or None if the known layout isn't found."""
    section_match = _TABLE_SECTION.search(text)
    window = text[section_match.start():section_match.start() + 4000] if section_match else text
    rows: list[dict] = []
    last_r_value: str | None = None
    last_thickness: str | None = None
    for line in window.splitlines():
        full = _BATT_FULL_ROW.match(line)
        if full:
            r_value, thickness, width, length, batts, m2, coverage, bale, code = full.groups()
            last_r_value, last_thickness = r_value, thickness
            rows.append({f"c{i}": v for i, v in enumerate([r_value, thickness, width, length, batts, m2, coverage, bale, code])})
            continue
        cont = _BATT_CONT_ROW.match(line)
        if cont and last_r_value is not None:
            width, length, batts, m2, coverage, bale, code = cont.groups()
            rows.append({f"c{i}": v for i, v in enumerate([last_r_value, last_thickness, width, length, batts, m2, coverage, bale, code])})
    if len(rows) < 2:
        return None
    return _BATT_TABLE_HEADERS, rows


_SKU_PDFS: dict[str, str] | None = None


def _sku_pdf_url(family_id: str, official: list[str]) -> str | None:
    """Official TDS PDF already recorded in the SKU catalogue for this family."""
    global _SKU_PDFS
    if _SKU_PDFS is None:
        _SKU_PDFS = {}
        import pandas as pd
        csv_path = ROOT / "data" / "processed" / "product_catalogue_skus.csv"
        if csv_path.exists():
            df = pd.read_csv(csv_path).fillna("")
            for _, row in df.iterrows():
                url = str(row["tds_url"]).strip()
                if url.lower().endswith(".pdf") and row["family_id"] not in _SKU_PDFS:
                    _SKU_PDFS[row["family_id"]] = url
    url = _SKU_PDFS.get(family_id)
    if url and domain_matches(domain_of(url), official):
        return url
    return None


def research_path(manufacturer_dir: str, slug: str) -> Path:
    return ROOT / "knowledge" / manufacturer_dir / RESEARCH_DIR_NAME / f"{slug}.json"


def fetch_pdf(url: str) -> Path | None:
    """Download a TDS document (PDF or DOCX) to the cache; return its path or
    None. Named fetch_pdf for historical reasons but handles both types —
    several manufacturers (e.g. Acoustica) only publish DOCX datasheets."""
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    ext = ".docx" if url.lower().split("?")[0].endswith(".docx") else ".pdf"
    name = hashlib_name(url) + ext
    target = CACHE_DIR / name
    if target.exists() and target.stat().st_size > 1000:
        return target
    try:
        response = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=30, stream=True)
        if response.status_code != 200:
            return None
        content = response.content
        if ext == ".pdf" and not content.startswith(b"%PDF"):
            return None
        if ext == ".docx" and not content.startswith(b"PK\x03\x04"):  # DOCX is a zip archive
            return None
        target.write_bytes(content)
        return target
    except Exception:
        return None


def hashlib_name(url: str) -> str:
    import hashlib
    return hashlib.sha1(url.encode()).hexdigest()[:16]


def pdf_text(path: Path, max_pages: int = 12) -> str:
    if path.suffix.lower() == ".docx":
        return docx_text(path)
    try:
        reader = pypdf.PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages[:max_pages]]
        return "\n".join(pages)
    except Exception:
        return ""


def docx_text(path: Path) -> str:
    try:
        import docx
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    except Exception:
        return ""


def search_tds_url(manufacturer: str, family_name: str) -> str | None:
    """Find a datasheet PDF on the manufacturer's own domain.

    Tries, in order: the site search engines (DuckDuckGo HTML, Bing), then a
    sitemap crawl of the official domain looking for PDF links whose path
    matches the family name. Search engines may be blocked in some
    environments; the sitemap crawl is the reliable fallback.
    """
    official = OFFICIAL_DOMAINS.get(manufacturer) or []
    if not official:
        return None
    domain = official[0]
    family_terms = [t for t in re.findall(r"[a-z0-9]+", family_name.casefold()) if len(t) > 3 and t != manufacturer.casefold()]

    def matches(url: str) -> bool:
        path = urllib.parse.urlparse(url).path.casefold()
        return bool(family_terms) and any(term in path for term in family_terms)

    for engine_url in (
        "https://html.duckduckgo.com/html/?q=" + urllib.parse.quote(f"site:{domain} {family_name} filetype:pdf"),
        "https://www.bing.com/search?q=" + urllib.parse.quote(f"site:{domain} {family_name} filetype:pdf"),
    ):
        try:
            response = requests.get(engine_url, headers={"User-Agent": USER_AGENT}, timeout=20)
            for link in re.findall(r'href="([^"]+\.pdf[^"]*)"', response.text, re.I):
                link = "https:" + link if link.startswith("//") else link
                if domain_matches(domain_of(link), official) and matches(link):
                    return link
        except Exception:
            continue

    # sitemap fallback: walk the official domain's sitemap(s) for product/PDF URLs
    try:
        return _sitemap_pdf(domain, family_terms, matches)
    except Exception:
        return None


def _sitemap_pdf(domain: str, family_terms: list[str], matches) -> str | None:
    seen: set[str] = set()
    queue = [f"https://{domain}/sitemap.xml", f"https://www.{domain}/sitemap.xml"]
    product_pages: list[str] = []
    for _ in range(6):  # bounded crawl
        if not queue:
            break
        sm = queue.pop(0)
        if sm in seen:
            continue
        seen.add(sm)
        try:
            xml = requests.get(sm, headers={"User-Agent": USER_AGENT}, timeout=20).text
        except Exception:
            continue
        locs = re.findall(r"<loc>([^<]+)</loc>", xml)
        for loc in locs:
            low = loc.casefold()
            if low.endswith(".pdf") and matches(loc):
                return loc
            if low.endswith(".xml"):
                queue.append(loc)
            elif matches(loc):
                product_pages.append(loc)
    # fetch a few candidate product pages and look for a linked PDF
    for page in product_pages[:5]:
        try:
            html = requests.get(page, headers={"User-Agent": USER_AGENT}, timeout=20).text
        except Exception:
            continue
        for link in re.findall(r'href="([^"]+\.pdf[^"]*)"', html, re.I):
            if link.startswith("/"):
                link = f"https://{domain}" + link
            if domain_matches(domain_of(link), [domain]):
                return link
    return None


def _validate_range(spec: dict) -> None:
    """Drop range/range_headers in place if the model produced a misaligned
    table (e.g. merged source cells it failed to repeat per row). A table we
    can't verify column-by-column must not be rendered at all."""
    headers = spec.get("range_headers")
    rows = spec.get("range")
    if not isinstance(headers, list) or not isinstance(rows, list) or not headers or not rows:
        spec["range_headers"], spec["range"] = [], []
        return
    expected_keys = {f"c{i}" for i in range(len(headers))}
    if any(not isinstance(row, dict) or set(row.keys()) != expected_keys for row in rows):
        spec["range_headers"], spec["range"] = [], []
        spec["range_extraction_status"] = "inconsistent_columns_dropped"


def extract_spec(text: str, need_range: bool = True) -> dict | None:
    """Ask the local Ollama model to structure the TDS text into JSON.

    `need_range` is False when parse_variant_table() already found the size
    table deterministically, so the model isn't asked to reproduce it (smaller
    prompt/output, and no risk of it misaligning a merged-cell table).
    """
    if not llm_client.ollama_available():
        return None
    trimmed = text[:16000]  # leave headroom for the JSON reply in context
    for attempt in range(2):
        raw = _generate_json(trimmed, need_range)
        if raw:
            spec = _parse_spec(raw)
            if spec is not None:
                if need_range:
                    _validate_range(spec)
                else:
                    spec.setdefault("range_headers", [])
                    spec.setdefault("range", [])
                return spec
        time.sleep(2 * (attempt + 1))  # back off; cold model loads can drop a call
    return None


def _adaptive_budget(text: str, need_range: bool) -> tuple[int, int]:
    """Size num_predict/num_ctx to the actual job instead of a fixed worst-case
    budget, so small datasheets (or ones where regex already got the range
    table) run several times faster on CPU-only local inference."""
    num_predict = 3200 if need_range else 1400
    estimated = len(text) // 4 + num_predict + 256
    num_ctx = max(4096, min(16384, 1 << (estimated - 1).bit_length()))
    return num_predict, num_ctx


def _generate_json(text: str, need_range: bool = True) -> str | None:
    """Direct Ollama call tuned for deterministic JSON: temperature 0 and a
    budget sized to the job, unlike the chat phrasing defaults in llm_client."""
    prompt = BASE_EXTRACT_PROMPT.replace("{range_fields}", RANGE_FIELDS_PROMPT if need_range else "")
    num_predict, num_ctx = _adaptive_budget(text, need_range)
    payload = {
        "model": EXTRACT_MODEL,
        "messages": [
            {"role": "system", "content": "You extract structured JSON from technical datasheets. Output only valid JSON, no markdown fences, no commentary."},
            {"role": "user", "content": prompt + text},
        ],
        "stream": False,
        "think": False,
        "keep_alive": "30m",
        "options": {"temperature": 0, "num_predict": num_predict, "num_ctx": num_ctx},
    }
    request = urllib.request.Request(
        f"{llm_client.OLLAMA_HOST}/api/chat",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=90) as response:
            data = json.loads(response.read().decode("utf-8"))
    except Exception:
        return None
    return (data.get("message") or {}).get("content", "").strip() or None


def _parse_spec(raw: str) -> dict | None:
    match = re.search(r"\{.*\}", raw, re.S)
    if not match:
        return None
    candidate = match.group(0)
    try:
        return json.loads(candidate)
    except json.JSONDecodeError:
        pass
    # truncated output: try progressively closing open arrays/object
    for closers in ('"]}', '"}]', "]}", "}", "]}"):
        try:
            return json.loads(candidate + closers)
        except json.JSONDecodeError:
            continue
    return None


_IDENTITY_STOPWORDS = {
    "insulation", "product", "products", "range", "system", "systems", "board", "panel",
    "the", "and", "for", "with", "series",
}


def _significant_terms(family_name: str, manufacturer_dir: str) -> list[str]:
    """Distinctive words from a family name, used to verify a datasheet is
    actually about this family and not a cross-linked/wrong-product PDF."""
    stop = _IDENTITY_STOPWORDS | {manufacturer_dir.casefold()}
    words = re.findall(r"[a-z0-9]+", family_name.casefold())
    return [w for w in words if len(w) > 2 and w not in stop]


def verify_family_identity(family_name: str, manufacturer_dir: str, text: str) -> bool:
    """True if the extracted PDF text plausibly matches this family: a
    domain being official (audit_datasheet_links.py) does not guarantee the
    specific linked document is for the right product — spreadsheet-inherited
    links have pointed multiple families at the same generic TDS before."""
    terms = _significant_terms(family_name, manufacturer_dir)
    if not terms:
        return True  # nothing distinctive to check (e.g. a bare "Accessory" family)
    haystack = text.casefold()
    hits = sum(1 for term in terms if term in haystack)
    # a short, distinctive name (the common case) must match in full — a
    # majority vote lets one generic word (e.g. "acoustic") false-positive a
    # completely different product's datasheet.
    required = len(terms) if len(terms) <= 3 else max(3, round(len(terms) * 0.8))
    return hits >= required


def process_family(manufacturer_dir: str, family: dict, delay: float = 1.0, refresh: bool = False) -> str:
    slug = slugify(family["name"])
    out_path = research_path(manufacturer_dir, slug)
    previous_pdf_url: str | None = None
    if out_path.exists():
        try:
            previous = json.loads(out_path.read_text(encoding="utf-8"))
            if previous.get("status") == "ok" and not refresh:
                return "cached"
            previous_pdf_url = previous.get("datasheet_pdf_url")
        except (json.JSONDecodeError, OSError):
            pass  # corrupt file: reprocess

    tds_url = (family.get("source_url") or "").strip()
    official = OFFICIAL_DOMAINS.get(manufacturer_dir) or []
    on_official = tds_url and domain_matches(domain_of(tds_url), official)

    # candidate PDF URLs in preference order. A domain being official
    # (audit_datasheet_links.py) does not prove the specific document is for
    # this family, so every candidate is fetched and identity-checked below
    # rather than trusting the first one that resolves.
    candidates: list[str] = []
    sku_pdf = _sku_pdf_url(family["family_id"], official)
    if sku_pdf:
        candidates.append(sku_pdf)
    if tds_url.lower().split("?")[0].endswith((".pdf", ".docx")) and on_official and tds_url not in candidates:
        candidates.append(tds_url)
    if previous_pdf_url and previous_pdf_url not in candidates:
        candidates.append(previous_pdf_url)

    tried: list[str] = []
    text = pdf_url = None
    for candidate_url in candidates:
        tried.append(candidate_url)
        pdf_path = fetch_pdf(candidate_url)
        if not pdf_path:
            continue
        candidate_text = pdf_text(pdf_path)
        if len(candidate_text) < 200:
            continue
        if verify_family_identity(family["name"], manufacturer_dir, candidate_text):
            text, pdf_url = candidate_text, candidate_url
            break

    if text is None:
        searched_url = search_tds_url(manufacturer_dir, family["name"])
        time.sleep(delay)  # be polite to search
        if searched_url and searched_url not in tried:
            tried.append(searched_url)
            pdf_path = fetch_pdf(searched_url)
            if pdf_path:
                candidate_text = pdf_text(pdf_path)
                if len(candidate_text) >= 200 and verify_family_identity(family["name"], manufacturer_dir, candidate_text):
                    text, pdf_url = candidate_text, searched_url

    if text is None:
        if not tried:
            _write(out_path, family, None, None, None, "no_pdf_found")
            return "no_pdf"
        # documents were found but none were verifiably about this family
        _write(out_path, family, tried[-1], None, None, "identity_mismatch")
        return "identity_mismatch"

    variant_table = parse_variant_table(text)
    spec = extract_spec(text, need_range=variant_table is None)
    if spec is None:
        _write(out_path, family, pdf_url, None, None, "extract_failed")
        return "extract_failed"
    if variant_table is not None:
        headers, rows = variant_table
        spec["range_headers"], spec["range"] = headers, rows
        spec["range_extraction_status"] = "regex_parsed"

    _write(out_path, family, pdf_url, spec, text[:2000], "ok")
    return "ok"


def _write(out_path: Path, family: dict, pdf_url: str | None, spec: dict | None, excerpt: str | None, status: str) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps({
        "family_id": family["family_id"],
        "family_name": family["name"],
        "datasheet_pdf_url": pdf_url,
        "status": status,
        "researched_at": time.strftime("%Y-%m-%d"),
        "spec": spec,
        "source_excerpt": excerpt,
    }, indent=2, ensure_ascii=False), encoding="utf-8")


def status_report() -> None:
    counts: dict[str, int] = {}
    for path in sorted(ROOT.glob("knowledge/*/research/*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        counts[data["status"]] = counts.get(data["status"], 0) + 1
    total = sum(counts.values())
    print(f"researched families: {total}")
    for status, count in sorted(counts.items(), key=lambda kv: -kv[1]):
        print(f"  {status:<16} {count}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--only", help="one manufacturer directory name")
    parser.add_argument("--family", help="one family_id, for reprocessing a single family")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--delay", type=float, default=1.0)
    parser.add_argument("--refresh", action="store_true", help="reprocess families already marked ok (e.g. after a schema/prompt change)")
    parser.add_argument("--status", action="store_true")
    args = parser.parse_args()

    if args.status:
        status_report()
        return

    if not llm_client.ollama_available():
        print("WARNING: local Ollama not detected - extraction will fail. Start 'ollama serve' first.")

    # A single family (a slow/hanging search or fetch) must never stall an
    # unattended multi-hour run: bound each one with a hard wall-clock timeout.
    # A fresh single-use executor per family means an abandoned/hung thread
    # (Python threads can't be force-killed) never blocks the next family.
    FAMILY_TIMEOUT_SECONDS = 240

    def process_with_timeout(manufacturer_dir: str, family: dict) -> str:
        executor = ThreadPoolExecutor(max_workers=1)
        future = executor.submit(process_family, manufacturer_dir, family, args.delay, args.refresh)
        try:
            return future.result(timeout=FAMILY_TIMEOUT_SECONDS)
        except FutureTimeoutError:
            slug = slugify(family["name"])
            _write(research_path(manufacturer_dir, slug), family, None, None, None, "timeout")
            return "timeout"
        finally:
            executor.shutdown(wait=False)

    done = 0
    tally: dict[str, int] = {}
    for path in sorted(ROOT.glob("knowledge/*/families.json")):
        manufacturer_dir = path.parent.name
        if args.only and manufacturer_dir.casefold() != args.only.casefold():
            continue
        data = json.loads(path.read_text(encoding="utf-8"))
        for family in data["families"]:
            family.setdefault("manufacturer", manufacturer_dir.title())
            if args.family and family["family_id"] != args.family:
                continue
            if args.limit and done >= args.limit:
                break
            result = process_with_timeout(manufacturer_dir, family)
            tally[result] = tally.get(result, 0) + 1
            done += 1
            if result != "cached":
                print(f"[{done}] {family['name'][:50]:<52} {result}")
        if args.limit and done >= args.limit:
            break

    print("\ntally:", tally)


if __name__ == "__main__":
    main()
