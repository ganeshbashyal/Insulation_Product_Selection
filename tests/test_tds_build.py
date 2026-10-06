"""Desktop ingestion uses exact provenance, explicit gaps and local-only models."""
import json
import io
from pathlib import Path
from types import SimpleNamespace
import zipfile

from openpyxl import Workbook
import pytest

import tds_build as tds


@pytest.fixture
def setup(tmp_path, monkeypatch):
    root = tmp_path / "repo"
    root.mkdir()
    cache = tmp_path / "cache"
    cache.mkdir()
    index = SimpleNamespace(root=root, families={"TEST": {"name": "Test Family", "manufacturer": "Test"}},
                            detail=lambda key: {"sources": {"documents": []}})
    monkeypatch.setattr(tds, "build_preview", lambda idx: ({"preview_id": "AUTHORING"}, {}))
    workbook = tmp_path / "links.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "TDS needed"
    sheet.append(["Manufacturer", "Family Name", "Family ID", "Status", "Existing URL", "SUPPLIED TDS URL <- fill in"])
    sheet.append(["Test", "Test Family", "TEST", "ok", "", "https://example.org/source.pdf"])
    sheet.append(["Test", "Test Family", "TEST", "", "", "https://example.org/shared.pdf"])
    sheet["F3"].hyperlink = "https://example.org/different.pdf"
    sheet.append(["Test", "Unknown", "UNKNOWN", "", "", "https://example.org/other.pdf"])
    sheet.append(["Test", "Test Family", "", "", "https://example.org/page", "=E5"])
    second = book.create_sheet("Sheet2")
    second.append(["Manufacturer", "Family Name", "SUPPLIED TDS URL"])
    second.append(["Test", "Test Family", "https://example.org/source.pdf"])
    book.save(workbook)
    return tds.Build(root, cache), index, workbook


def test_preview_reads_all_sheets_conflicts_formula_no_writes(setup):
    build, index, workbook = setup
    before = workbook.read_bytes()
    preview = build.preview(workbook, index)
    assert preview["counts"] == {"supplied_rows": 5, "held_rows": 2, "eligible_urls": 2, "local_files": 0}
    assert preview["rows"][1]["holds"] == ["display_hyperlink_conflict"]
    assert preview["rows"][2]["holds"] == ["unknown_or_unresolved_family"]
    assert preview["rows"][3]["formula"] == "=E5"
    assert preview["rows"][4]["family_id"] == "TEST"
    assert not build.base.exists() and workbook.read_bytes() == before


def test_archive_deduplicates_content_and_extracts_docx_without_models(setup):
    build, index, workbook = setup
    doc = build.cache / "unknown.docx"
    with zipfile.ZipFile(doc, "w") as z:
        z.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Glass wool</w:t></w:r></w:p></w:body></w:document>')
    original = doc.read_bytes()
    preview = build.preview(workbook, index)
    build.start(preview, preview["build_id"])
    assert build.status(preview["build_id"])["status"]["local"] == {"archived": 1}
    build.extract(preview["build_id"])
    chunks = list(build.chunks(preview["build_id"]))
    assert chunks[0]["text"] == "Glass wool"
    assert chunks[0]["locator"].startswith("paragraph")
    build.start(preview, preview["build_id"])
    assert len(list(build.path("originals").iterdir())) == 1
    assert doc.read_bytes() == original
    source = next(iter(build.sources(preview["build_id"]).values()))
    assert source["families"] == [] and "unknown" in source["rights"]
    cache_receipt = next(build.path("extraction").glob("*.json"))
    corrupt = json.loads(cache_receipt.read_text())
    corrupt["pages"][0]["text"] = "altered"
    cache_receipt.write_text(json.dumps(corrupt))
    with pytest.raises(ValueError, match="checksum"):
        list(build.chunks(preview["build_id"]))


def test_progress_sheet_uses_exact_embedded_ids_and_holds_conflicts(setup):
    build, index, workbook = setup
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet3"
    sheet.append(["#", "Family", "State", "TDS link"])
    sheet.append([1, "Test Family (TEST)", "needs_source", "https://example.org/new.pdf"])
    sheet.append([2, "Test Family (UNKNOWN)", "needs_source", "https://example.org/unknown.pdf"])
    sheet.append([3, "Test Family (TEST)", "needs_source", "https://example.org/display.pdf"])
    sheet["D4"].hyperlink = "https://example.org/different.pdf"
    book.save(workbook)
    preview = build.preview(workbook, index)
    assert preview["counts"]["supplied_rows"] == 3
    assert preview["counts"]["eligible_urls"] == 1
    assert preview["rows"][0]["family_id"] == "TEST"
    assert preview["rows"][0]["identity_method"] == "supplied_embedded_id"
    assert preview["rows"][1]["holds"] == ["unknown_or_unresolved_family"]
    assert preview["rows"][2]["holds"] == ["display_hyperlink_conflict"]


def test_additive_build_retains_archives_receipts_and_model_history(setup):
    build, index, workbook = setup
    first = build.preview(workbook, index)
    build.start(first, first["build_id"])
    directory, _ = build.job(first["build_id"])
    doc = build.cache / "source.docx"
    with zipfile.ZipFile(doc, "w") as z:
        z.writestr("word/document.xml", "<document/>")
    archived = build.archive_file(doc)
    url = "https://example.org/source.pdf"
    tds.save(directory / "downloads" / (tds.checksum(url)[:24] + ".json"),
             {**archived, "url": url, "status": "archived"})
    tds.save(directory / "model" / "historic.json", {"status": "failed", "raw": "preserved"})
    tds.save(directory / "splits" / "historic.json", {"input": "preserved"})
    doc.unlink()
    second = build.preview(workbook, index, extend_from=first["build_id"])
    assert second["extends_build"] == first["build_id"]
    assert second["local_files"][0]["families"] == ["TEST"]
    build.start(second, second["build_id"])
    newer, _ = build.job(second["build_id"])
    assert build.sources(second["build_id"])[archived["sha256"]]["families"] == ["TEST"]
    assert (newer / "model" / "historic.json").read_bytes() == (directory / "model" / "historic.json").read_bytes()
    assert (newer / "splits" / "historic.json").exists()
    assert build.status(second["build_id"])["status"]["downloads"] == {"archived": 1}
    Path(archived["original"]).write_bytes(b"changed")
    with pytest.raises(ValueError, match="Previous archived source"):
        build.preview(workbook, index, extend_from=first["build_id"])


def test_manifest_stale_workbook_and_safe_paths(setup):
    build, index, workbook = setup
    preview = build.preview(workbook, index)
    with pytest.raises(ValueError, match="confirmation"):
        build.start(preview, "wrong")
    with pytest.raises(ValueError, match="escapes"):
        build.path("..", "..", "outside")
    workbook.write_bytes(b"changed")
    with pytest.raises(ValueError, match="changed"):
        build.start(preview, preview["build_id"])


@pytest.mark.parametrize("url", ["http://example.org/a.pdf", "file:///x", "https://user:pw@example.org/a",
                                "https://example.org:22/a", "https://example.org/\nfoo"])
def test_invalid_public_link_shapes(url):
    with pytest.raises(ValueError):
        tds.validate_url_shape(url)


def test_private_dns_and_cross_host_redirect_held(monkeypatch):
    monkeypatch.setattr(tds.socket, "getaddrinfo", lambda *a, **k: [(2, 1, 6, "", ("127.0.0.1", 443))])
    with pytest.raises(ValueError, match="Private"):
        tds.public_url("https://example.org/a.pdf")
    monkeypatch.setattr(tds, "public_url", lambda u: None)
    req = tds.urllib.request.Request("https://example.org/a")
    with pytest.raises(ValueError, match="Cross-host"):
        tds.CheckedRedirect().redirect_request(req, None, 302, "", {}, "https://other.org/a")


def test_exact_quotes_are_not_model_approvals():
    assert tds.validate_candidates('{"candidates":[{"quote":"Glass wool","topic":"material"}]}', "Glass wool")
    for raw in ('{"candidates":[{"quote":"R9","topic":"performance"}]}',
                '{"candidates":[{"quote":"Glass wool","topic":"approved"}]}',
                '{"candidates":[],"approval":true}'):
        with pytest.raises(ValueError):
            tds.validate_candidates(raw, "Glass wool")
    raw = '{"candidates":[{"start_line":1,"end_line":2,"topic":"material"}]}'
    assert tds.validate_candidates(raw, "Glass wool\nUnnormalised  OCR")[0]["quote"] == "Glass wool\nUnnormalised  OCR"
    with pytest.raises(ValueError, match="out-of-range"):
        tds.validate_candidates('{"candidates":[{"start_line":1,"end_line":9,"topic":"material"}]}', "One line")


class Client:
    def __init__(self, invalid=False):
        self.calls = []
        self.invalid = invalid

    def request(self, endpoint, body=None):
        self.calls.append((endpoint, body))
        if endpoint == "/api/tags":
            return {"models": [{"name": "llama3.1:8b"}]}
        if endpoint == "/api/ps":
            return {"models": []}
        if endpoint == "/api/generate":
            return {}
        return {"message": {"content": '{"candidates":[{"quote":"invented","topic":"material"}]}' if self.invalid
                            else '{"candidates":[]}'}}


def test_model_uses_two_threads_unloads_and_records_failures(setup):
    build, _, _ = setup
    client = Client()
    result = build.model_call(client, {"model": "llama3.1:8b", "text": "Glass wool"})
    assert result["status"] == "validated"
    assert client.calls[-1][0] == "/api/generate"
    assert next(body for endpoint, body in client.calls if endpoint == "/api/chat")["options"]["num_thread"] == 2
    result = build.model_call(Client(invalid=True), {"model": "llama3.1:8b", "text": "Glass wool"})
    assert result["status"] == "failed" and result["raw"] and "ungrounded" in result["error"]
    schema = next(body for endpoint, body in client.calls if endpoint == "/api/chat")["format"]
    assert schema["properties"]["candidates"]["items"]["anyOf"][0]["required"] == ["start_line", "end_line", "topic"]


def test_download_failure_saved_and_not_repeated(setup, monkeypatch):
    build, index, workbook = setup
    preview = build.preview(workbook, index)
    build.start(preview, preview["build_id"])
    calls = []
    def fail(url, path):
        calls.append(url)
        raise ValueError("HTML is not PDF")
    monkeypatch.setattr(tds, "download_document", fail)
    assert build.download(preview["build_id"], 2)["status"]["downloads"] == {"failed": 2}
    assert build.download(preview["build_id"], 2)["attempted_this_run"] == 0
    assert len(calls) == 2


def test_download_bytes_validation_and_size_budget(tmp_path, monkeypatch):
    monkeypatch.setattr(tds, "public_url", lambda url: None)
    class Reply(io.BytesIO):
        url = "https://example.org/file.pdf"
        status = 200
        headers = {}
    class Opener:
        content = b"%PDF-1.7\nsynthetic"
        def open(self, *args, **kwargs):
            return Reply(self.content)
    opener = Opener()
    monkeypatch.setattr(tds.urllib.request, "build_opener", lambda *a: opener)
    target = tmp_path / "download.part"
    result = tds.download_document("https://example.org/file.pdf", target)
    assert result["bytes"] == len(opener.content) and result["kind"] == ".pdf"
    target.unlink()
    opener.content = b"<html>Not a PDF</html>"
    with pytest.raises(ValueError, match="Not a supported"):
        tds.download_document("https://example.org/file.pdf", target)
    assert not target.exists()
    monkeypatch.setattr(tds, "MAX_BYTES", 3)
    opener.content = b"%PDF-oversized"
    with pytest.raises(ValueError, match="budget"):
        tds.download_document("https://example.org/file.pdf", target)
    assert not target.exists()


def test_private_packs_preserve_retained_data_and_rights_gate(setup, monkeypatch):
    build, index, workbook = setup
    row = {"field": "material", "value": "old value", "origin": "legacy", "locator": "spec",
           "status": "retained_not_reviewed"}
    pack = {"family_id": "TEST", "name": "Synthetic",
            "dossier": {"groups": {"material": [row]}, "alternatives": [],
                        "retained": {"research": {"unknown": [1, None]}, "commercial_rows": []},
                        "provenance_state": "legacy_supplied_original_unavailable"}}
    stem = tds.hashlib.sha256(b"TEST").hexdigest()[:20]
    authoring = {"preview_id": "AUTHORING", "pack": pack}
    monkeypatch.setattr(tds, "build_preview", lambda idx: (
        {"preview_id": authoring["preview_id"], "families": {"TEST": {"file": stem}}},
        {"TEST": authoring["pack"]}))
    preview = build.preview(workbook, index)
    build.start(preview, preview["build_id"])
    result = build.packs(preview["build_id"], index)
    assert result["family_count"] == 1
    assert result["status"]["status"]["pending_urls"] == 2
    assert build.family("TEST", index)["pack"]["dossier"] == pack["dossier"]
    assert "old value" in build.family("TEST", index)["readable"]
    again = build.packs(preview["build_id"], index)
    assert again["path"] == result["path"]

    updated_pack = {**pack, "dossier": {**pack["dossier"], "groups": {
        "material": [{**row, "value": "updated value"}]}}}
    authoring.update(preview_id="UPDATED_AUTHORING", pack=updated_pack)
    assert build.family("TEST", index)["authoring_stale"]
    refreshed = build.packs(preview["build_id"], index)
    assert refreshed["authoring_preview"] == "UPDATED_AUTHORING"
    current = build.family("TEST", index)
    assert not current["authoring_stale"]
    assert current["pack"]["dossier"]["groups"]["material"][0]["value"] == "updated value"
    report_result = build.report(preview["build_id"], index)
    report = json.loads(build.path("reports", "completion.json").read_text())
    assert report_result["summary"] == {"draft_with_source_gap": 1}
    assert report["build_authoring_changed"]
    assert report["families"][0]["draft_authoring_preview"] == "UPDATED_AUTHORING"
    assert not report["families"][0]["authoring_stale"]
    assert build.packs(preview["build_id"], index)["path"] == refreshed["path"]

    metadata = json.loads(build.path("training_candidates", preview["build_id"] + ".json").read_text())
    assert metadata["eligible_content_records"] == 0 and not metadata["fine_tuning"]
    assert (build.root / "data" / "local" / "fresh_tds_cache.json").exists()


def test_desktop_backup_copies_new_independent_directory(setup, tmp_path):
    build, index, workbook = setup
    preview = build.preview(workbook, index)
    build.start(preview, preview["build_id"])
    result = build.backup(tmp_path / "independent-backup")
    assert result["verified"] and result["files"] >= 2
    with pytest.raises(ValueError, match="independent"):
        build.backup(tmp_path / "independent-backup")


def test_family_packs_preserve_other_family_pointer(setup, monkeypatch):
    build, index, workbook = setup
    index.families["SECOND"] = {"name": "Second"}
    packs = {key: {"family_id": key, "name": key,
                  "dossier": {"groups": {}, "alternatives": [], "retained": {},
                              "provenance_state": "legacy_supplied_original_unavailable"}}
             for key in index.families}
    families = {key: {"file": tds.hashlib.sha256(key.encode()).hexdigest()[:20]} for key in packs}
    monkeypatch.setattr(tds, "build_preview", lambda idx: (
        {"preview_id": "AUTHORING", "families": families}, packs))
    preview = build.preview(workbook, index)
    build.start(preview, preview["build_id"])
    first = build.packs(preview["build_id"], index, family_id="TEST")
    assert first["family_count"] == 1
    assert build.family("SECOND", index)["state"] == "not_generated"
    second = build.packs(preview["build_id"], index, family_id="SECOND")
    assert second["family_count"] == 1
    assert build.family("TEST", index)["pack"]["family_id"] == "TEST"
    assert build.family("SECOND", index)["pack"]["family_id"] == "SECOND"


def test_family_selector_mandatory_in_cli():
    from scripts.fresh_tds_build import main
    with pytest.raises(SystemExit) as exc:
        main(["--cache", "Cache", "run", "a" * 64])
    assert exc.value.code == 2


def test_completion_report_alphabetical_and_no_false_completion(setup):
    build, index, workbook = setup
    index.families["FIRST"] = {"name": "First accessory", "manufacturer": "Acoustica"}
    preview = build.preview(workbook, index)
    build.start(preview, preview["build_id"])
    result = build.report(preview["build_id"], index)
    assert result["first_family"]["family_id"] == "FIRST"
    assert result["summary"] == {"needs_source": 2}
    report = json.loads(build.path("reports", "completion.json").read_text())
    assert all(not row["draft_generated"] and row["human_approval"] == "not_granted_by_build"
               for row in report["families"])
    assert "Chunks validated / total" in Path(result["path"]).read_text()
    report_text = Path(result["path"]).read_text()
    assert "Chat with the bot about one family at a time" in report_text
    assert "[Chat with bot](http://127.0.0.1:8001/chat?family_id=FIRST)" in report_text
    assert "exact family preselected as conversation context" in report_text
    assert "No supplemental link workbook was supplied" in Path(result["path"]).read_text()
    assert report["supplemental_link_status"] == "not_supplied"
    assert report["supplemental_link_workbook"] is None
    assert Path(result["csv"]).read_text().startswith("order,family_id,")
    assert "Missing primary source" in Path(result["path"]).read_text()


def test_completion_report_includes_unverified_supplemental_tds_links(setup):
    build, index, workbook = setup
    preview = build.preview(workbook, index)
    build.start(preview, preview["build_id"])
    supplemental = {
        "workbook": str(workbook),
        "sha256": "a" * 64,
        "rows": [
            {"family_id": "TEST", "family_name": "Test Family", "sheet": "Links",
             "row": 7, "url": "https://example.org/test-tds.pdf", "holds": []},
            {"family_id": "TEST", "family_name": "Test Family", "sheet": "Links",
             "row": 8, "url": "javascript:alert(1)", "holds": ["invalid_or_unresolved_url"]},
        ],
    }
    sku_inventory = {
        "source_workbook": "Staff release.xlsm",
        "source_sha256": "b" * 64,
        "baseline_location": "Melbourne",
        "active_product_master_skus": 2,
        "inactive_product_master_skus_excluded": 1,
        "price_values_included": False,
        "sales_location_summary": {
            "Melbourne": {"sku_count": 1, "melbourne_skus_not_listed": 0,
                          "state_only_skus": 0, "field_differences": {},
                          "sell_price_variances": 0},
            "Brisbane": {"sku_count": 2, "melbourne_skus_not_listed": 0,
                         "state_only_skus": 1, "field_differences": {"MOQ": 1},
                         "sell_price_variances": 1},
        },
        "products": [
            {"our_sku": "SKU-1", "family_id": "TEST", "candidate_family_id": "TEST",
             "family_mapping_status": "deterministic_high_unreviewed",
             "melbourne_baseline": True, "state_presence": ["Melbourne", "Brisbane"],
             "state_differences": {"Melbourne": [], "Brisbane": ["MOQ", "sell_price_varies"]},
             "price_statuses": {"Melbourne": "numeric", "Brisbane": "poa"}},
            {"our_sku": "SKU-2", "family_id": None, "candidate_family_id": None,
             "family_mapping_status": "unmapped_review_required",
             "family_codes": [], "product_name": "Unmapped product", "manufacturer": "Maker",
             "melbourne_baseline": False, "state_presence": ["Brisbane"],
             "state_differences": {"Brisbane": []}, "price_statuses": {"Brisbane": "numeric"}},
        ],
    }

    result = build.report(
        preview["build_id"], index, supplemental_links=supplemental,
        supplemental_skus=sku_inventory,
    )
    markdown = Path(result["path"]).read_text()
    report = json.loads(Path(result["json"]).read_text())

    assert "Supplemental workbook TDS links (unverified)" in markdown
    assert "[open supplied link](<https://example.org/test-tds.pdf>)" in markdown
    assert "`javascript:alert(1)`" in markdown
    assert "[open supplied link](<javascript:alert(1)>)" not in markdown
    assert report["supplemental_link_workbook"]["sha256"] == "a" * 64
    assert report["supplemental_link_status"] == "supplied_unverified"
    assert len(report["supplemental_tds_links"]) == 2
    assert report["summary"] == {"needs_source": 1}
    assert "Staff-release SKU inventory (unreviewed)" in markdown
    assert "SKU-1" in markdown and "SKU-2" in markdown
    assert "Sell-price variance" in markdown or "sell-price variance" in markdown
    assert "$" not in markdown
    assert report["supplemental_sku_inventory"]["source_sha256"] == "b" * 64
    assert report["staff_release_family_skus"][0]["products"][0]["our_sku"] == "SKU-1"
    assert report["staff_release_unassigned_skus"][0]["our_sku"] == "SKU-2"


def test_single_family_worker_budget_and_no_advance(monkeypatch, capsys, tmp_path):
    from scripts import fresh_tds_build as cli
    calls = []
    class FakeBuild:
        def __init__(self, root, cache):
            self.used = 0
        def model_batch(self, build_id, limit, family_id):
            assert family_id == "TEST"
            calls.append(limit)
            self.used += limit
            return {"calls_this_run": limit}
        def packs(self, build_id, index, family_id):
            assert family_id == "TEST"
            return {"processing_coverage": {"pending_chunks": 100 - self.used}}
        def report(self, build_id, index):
            return {"path": "report.md"}
        def sources(self, build_id, family_id):
            return {"source": {}}
        def family(self, family_id, index):
            return {"pack": {"dossier": {}, "fresh_full_extraction": []}, "readable": ""}
    monkeypatch.setattr(cli, "Build", FakeBuild)
    monkeypatch.setattr(cli, "ResearchIndex", lambda *a: SimpleNamespace(
        families={"TEST": {}}, detail=lambda key: {"dossier": {}}))
    assert cli.main(["--cache", str(tmp_path), "work-family", "a" * 64,
                     "--family", "TEST", "--max-calls", "6"]) == 0
    assert calls == [5, 1]
    assert "Worker stops at this family" in capsys.readouterr().out


def test_docx_paragraph_grouping_keeps_all_nonblank_text_and_locators():
    extraction = {"pages": [{"page": i, "text": f"original paragraph {i}"} for i in range(1, 120)]}
    units = tds.chunk_units(extraction, ".docx")
    assert len(units) < 5
    assert "\n".join(u["text"] for u in units) == "\n".join(p["text"] for p in extraction["pages"])
    assert units[0]["locator"].startswith("paragraphs 1-")
    assert all(len(u["text"]) <= 1000 for u in units)


def test_generation_schema_constrains_span_not_just_line_bounds():
    schema = tds.bounded_span_schema(43)
    choices = schema["properties"]["candidates"]["items"]["anyOf"]
    assert choices[22]["properties"]["start_line"]["const"] == 23
    assert choices[22]["properties"]["end_line"] == {"type": "integer", "minimum": 23, "maximum": 30}
    assert all(choice["required"] == ["start_line", "end_line", "topic"] for choice in choices)
    assert choices[-1]["properties"]["end_line"]["maximum"] == 43


def test_nonblank_schema_preserves_line_numbers_and_allows_no_candidates(setup):
    build, _, _ = setup
    schema = tds.bounded_span_schema(2, ["", "2400"])
    choices = schema["properties"]["candidates"]["items"]["anyOf"]
    assert choices[0]["properties"]["start_line"]["const"] == 1
    assert choices[0]["properties"]["end_line"]["minimum"] == 2
    assert choices[1]["properties"]["start_line"]["const"] == 2
    assert tds.bounded_span_schema(2, ["", " "])["properties"]["candidates"]["maxItems"] == 0
    assert tds.validate_candidates('{"candidates":[]}', "\n2400") == []
    job = {"text": "\n2400", "schema": tds.response_schema(2)}
    job["task_id"] = tds.checksum(job)
    revised = build.nonblank_task("a" * 64, job)
    assert revised["text"] == job["text"]
    assert revised["task_id"] != job["task_id"]
    tds.save(build.path("model_results", job["task_id"] + ".json"),
             {"input": job, "status": "validated", "raw": '{"candidates":[]}'})
    assert build.nonblank_task("a" * 64, job) == job


def test_overnight_chains_five_family_batches_with_total_budget(monkeypatch):
    from scripts import fresh_tds_build as cli
    calls = []
    def batch(build, index, build_id, budget, max_families):
        calls.append((budget, max_families))
        return {"families_this_worker": 5, "calls": min(3, budget), "report": {}}
    monkeypatch.setattr(cli, "work_alphabetical", batch)
    result = cli.work_overnight(None, None, "a" * 64, 7, 60)
    assert calls == [(7, 5), (4, 5), (1, 5)]
    assert result["calls"] == 7
    assert result["stop_reason"] == "call_budget"


def test_overnight_stops_on_exhaustion_batch_ceiling_and_failure(monkeypatch):
    from scripts import fresh_tds_build as cli
    monkeypatch.setattr(cli, "work_alphabetical", lambda *a: {
        "families_this_worker": 0, "calls": 0, "report": {}})
    assert cli.work_overnight(None, None, "a" * 64, 10, 60)["batches"] == 1
    monkeypatch.setattr(cli, "work_alphabetical", lambda *a: {
        "families_this_worker": 5, "calls": 0, "report": {}})
    assert cli.work_overnight(None, None, "a" * 64, 10, 2)["stop_reason"] == "batch_budget"
    attempts = []
    def failed(*args):
        attempts.append(1)
        raise ValueError("saved failure")
    monkeypatch.setattr(cli, "work_alphabetical", failed)
    with pytest.raises(ValueError, match="saved failure"):
        cli.work_overnight(None, None, "a" * 64, 10, 60)
    assert len(attempts) == 1


def test_alphabetical_worker_serial_budget_and_resume(tmp_path, monkeypatch, capsys):
    from scripts import fresh_tds_build as cli
    report = tmp_path / "progress.json"
    report.write_text(json.dumps({"families": [
        {"family_id": "A", "state": "draft_ready_for_human_review"},
        {"family_id": "B", "state": "model_pending"},
        {"family_id": "C", "state": "model_pending"}]}))
    class Fake:
        def __init__(self, *args):
            pass
        def report(self, *args):
            return {"json": str(report)}
    calls = []
    def one(build, index, bid, family, budget):
        calls.append((family, budget))
        return {"family_id": family, "calls": min(2, budget)}
    monkeypatch.setattr(cli, "Build", Fake)
    monkeypatch.setattr(cli, "ResearchIndex", lambda *a: object())
    monkeypatch.setattr(cli, "work_one", one)
    assert cli.main(["--cache", str(tmp_path), "work-alphabetical", "a" * 64,
                     "--max-calls", "3", "--max-families", "2"]) == 0
    assert calls == [("B", 3), ("C", 1)]
    assert '"calls": 3' in capsys.readouterr().out
