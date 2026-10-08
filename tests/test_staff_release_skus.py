import json

from openpyxl import Workbook
import pytest

from scripts.ingest_staff_release_skus import (
    SALES_FIELDS,
    apply_family_map_to_inventory,
    build_inventory,
)


def _write_workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Product_Master"
    sheet.append(["CORE IDENTITY"])
    master_headers = [
        "Our SKU", "Our Product Name", "Active?", "Category", "Material Type",
        "Product Use", "Manufacturer Name", "Spec Family Name", "Spec Material Type",
        "Spec ID",
    ]
    sheet.append(master_headers)
    sheet.append(["SKU-1", "Aircell Glareshield", "Yes", "Accessory", "Foil",
                  "General Installation", "Aircell", "", "", ""])
    sheet.append(["SKU-2", "Aircell Biscuits", "Yes", "Accessory", "Foil",
                  "General Installation", "Aircell", "", "", ""])
    sheet.append(["OLD-1", "Inactive line", "No", "Accessory", "Foil",
                  "General Installation", "Aircell", "", "", ""])

    migration = workbook.create_sheet("SKU_Migration_Register")
    migration.append(["SKU Migration Register"])
    migration.append(["Applied changes"])
    migration.append([])
    migration.append(["Status", "New / Canonical SKU", "Manufacturer", "Family"])
    migration.append(["RENAMED", "SKU-1", "Aircell", "GLARESHIELD"])
    migration.append(["RENAMED", "SKU-2", "Aircell", "BISCUITS"])

    for location in ("Melbourne", "Sydney", "Brisbane", "Adelaide", "Perth"):
        sales = workbook.create_sheet("Sales_" + location)
        sales.append([location + " sales view"])
        sales.append(["Staff-only"])
        sales.append(list(SALES_FIELDS))
        one = ["SKU-1", "Aircell Glareshield", "Accessory", "Foil",
               "General Installation", None, None, None, None, None, "Each",
               111.11 if location == "Melbourne" else ("POA" if location == "Sydney" else 222.22), None]
        if location == "Sydney":
            one[-1] = 4
        sales.append(one)
        if location == "Brisbane":
            sales.append(["SKU-2", "Aircell Biscuits", "Accessory", "Foil",
                          "General Installation", None, None, None, None, None, "Each",
                          333.33, None])
    workbook.save(path)


def test_staff_release_inventory_uses_melbourne_and_redacts_prices(tmp_path):
    source = tmp_path / "staff.xlsm"
    _write_workbook(source)

    inventory = build_inventory(source)
    products = {row["our_sku"]: row for row in inventory["products"]}

    assert len(products) == 2
    assert inventory["active_product_master_skus"] == 2
    assert inventory["inactive_product_master_skus_excluded"] == 1
    assert inventory["baseline_location"] == "Melbourne"
    assert products["SKU-1"]["melbourne_baseline"] is True
    assert products["SKU-2"]["melbourne_baseline"] is False
    assert products["SKU-2"]["state_presence"] == ["Brisbane"]
    assert inventory["sales_location_summary"]["Brisbane"]["state_only_skus"] == 1
    assert inventory["sales_location_summary"]["Sydney"]["sell_price_variances"] == 1
    assert inventory["sales_location_summary"]["Sydney"]["field_differences"] == {"MOQ": 1}
    encoded = json.dumps(inventory)
    assert "111.11" not in encoded and "222.22" not in encoded and "333.33" not in encoded
    assert inventory["price_values_included"] is False


def test_existing_inventory_applies_hash_bound_candidate_maps_conservatively(tmp_path):
    source = tmp_path / "staff.xlsm"
    _write_workbook(source)
    inventory = build_inventory(source)
    assert inventory["mapping_requests"]

    proposals = {}
    for request in inventory["mapping_requests"]:
        family_id = request["candidates"][0]["family_id"] if request["candidates"] else None
        proposals[request["group_id"]] = {
            "family_id": family_id, "confidence": 25, "reason": "Synthetic unreviewed candidate",
        }
    family_map = {
        "source_sha256": inventory["source_sha256"],
        "model": "synthetic-local-model",
        "groups": proposals,
    }
    applied = apply_family_map_to_inventory(inventory, family_map)
    assert applied["family_mapping_model"]["source_sha256"] == inventory["source_sha256"]
    assert applied["family_mapping_model"]["complete"] is True
    assert applied["family_mapping_model"]["review_groups_mapped"] == len(proposals)
    assert all(
        group["mapping_status"] == "local_model_candidate_unreviewed"
        for group in applied["groups"]
        if proposals.get(group["group_id"], {}).get("family_id")
        and not group["family_id"]
    )
    assert all(row["family_mapping_status"] != "approved" for row in applied["products"])
    assert applied["price_values_included"] is False

    wrong_hash = {**family_map, "source_sha256": "different"}
    with pytest.raises(ValueError, match="different workbook"):
        apply_family_map_to_inventory(build_inventory(source), wrong_hash)
    first_id = next(iter(proposals))
    partial = {**family_map, "groups": {first_id: proposals[first_id]}}
    partial_inventory = build_inventory(source)
    omitted_id = next(group_id for group_id in (
        row["group_id"] for row in partial_inventory["groups"]
    ) if group_id != first_id)
    before = next(row for row in partial_inventory["groups"] if row["group_id"] == omitted_id).copy()
    partial_applied = apply_family_map_to_inventory(partial_inventory, partial)
    omitted = next(row for row in partial_applied["groups"] if row["group_id"] == omitted_id)
    assert partial_applied["family_mapping_model"]["complete"] is False
    assert partial_applied["family_mapping_model"]["review_groups_mapped"] == 1
    assert omitted["candidate_family_id"] == before["candidate_family_id"]
    assert omitted["mapping_status"] == before["mapping_status"]
    crosswalk = {**family_map, "groups": {
        **proposals, first_id: {**proposals[first_id], "family_id": "NOT_A_CANDIDATE"},
    }}
    with pytest.raises(ValueError, match="non-candidate"):
        apply_family_map_to_inventory(build_inventory(source), crosswalk)
