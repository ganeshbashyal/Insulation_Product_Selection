import csv
import json

import pytest
from openpyxl import Workbook, load_workbook

from scripts import build_missing_tds_report as report
from scripts.knowledge_function_map import inventory


def test_report_covers_canonical_missing_research_and_ignores_ok_as_source(tmp_path, monkeypatch):
    knowledge = tmp_path / "knowledge"
    folder = knowledge / "test"
    folder.mkdir(parents=True)
    (folder / "families.json").write_text(json.dumps({"families": [
        {"family_id": "NO_RESEARCH", "name": "Wall Batt"},
        {"family_id": "OK_WITHOUT_PDF", "name": "Ceiling Batt"}]}))
    (knowledge / "_tds_manifest.json").write_text("{}")
    (folder / "research").mkdir()
    (folder / "research" / "test.json").write_text(json.dumps({"family_id": "OK_WITHOUT_PDF", "status": "ok"}))
    monkeypatch.setattr(report, "ROOT", tmp_path)
    rows = report._research_rows()
    assert {row["family_id"] for row in rows} == {"NO_RESEARCH", "OK_WITHOUT_PDF"}


def test_csv_keeps_manual_columns_and_resolved_requests(tmp_path):
    path = tmp_path / "requests.csv"
    path.write_text("family_id,supplied_tds_url,my_notes\nTEST,local supplied link,Keep me\nOLD,old link,Do not lose\n")
    report._write_csv(path, [{"family_id": "TEST", "supplied_tds_url": "", "missing_reason": "no_held_datasheet"}])
    with path.open() as handle:
        rows = {row["family_id"]: row for row in csv.DictReader(handle)}
    assert rows["TEST"]["my_notes"] == "Keep me"
    assert rows["TEST"]["supplied_tds_url"] == "local supplied link"
    assert rows["OLD"]["missing_reason"] == "historical_request_retained"


def test_workbook_keeps_manual_formula_and_extra_column(tmp_path):
    path = tmp_path / "requests.xlsx"
    book = Workbook()
    book.active.append(["family_id", "supplied_tds_url", "custom"])
    book.active.append(["TEST", "owner link", "=1+1"])
    book.save(path)
    assert report._write_xlsx(path, [{"family_id": "TEST", "supplied_tds_url": ""}])
    book = load_workbook(path, data_only=False)
    try:
        rows = list(book.active.values)
        assert rows[1][1:] == ("owner link", "=1+1")
    finally:
        book.close()


def test_duplicate_manual_rows_block_replacement(tmp_path):
    path = tmp_path / "requests.csv"
    original = "family_id,supplied_tds_url\nTEST,a\nTEST,b\n"
    path.write_text(original)
    with pytest.raises(ValueError, match="duplicate"):
        report._write_csv(path, [{"family_id": "TEST"}])
    assert path.read_text() == original


def test_function_map_retains_cli_and_decorated_uncertainty(tmp_path):
    (tmp_path / "module.py").write_text("@route.get('/x')\ndef detail():\n    return helper()\n\ndef helper():\n    return 1\n")
    (tmp_path / "scripts").mkdir()
    (tmp_path / "scripts" / "tool.py").write_text("def main():\n    pass\n")
    data = inventory(tmp_path)
    rows = {row["callable"]: row for row in data["functions"]}
    assert rows["main"]["entrypoint"] == "CLI"
    assert rows["detail"]["decorators"]
    assert rows["helper"]["syntactic_references"] == ["module.py:3"]
    assert rows["main"]["classification"] == "review_required"


def test_audit_updates_keep_previous_receipts_and_restage_idempotently(tmp_path):
    from local_source_review import retain_review
    output = tmp_path / "source_review.json"
    old = b'{"documents": {"old": "retained"}}'
    output.write_bytes(old)
    report = {"documents": {"new": "unreviewed"}, "auto_approval": False}
    first = retain_review(report, output)
    history = output.parent / "source_reviews"
    assert len(list(history.glob("*.json"))) == 2
    assert any(path.read_bytes() == old for path in history.glob("*.json"))
    assert retain_review(report, output)["review_id"] == first["review_id"]
    assert len(list(history.glob("*.json"))) == 2


def test_authoring_health_uses_shared_statuses_not_optional_legacy_stores(monkeypatch):
    import data_health
    import knowledge_service
    monkeypatch.setattr(knowledge_service.KnowledgeService, "validation", lambda self: {
        "family_count": 3, "families_with_linked_sources": 1, "sku_count": 2,
        "families_without_skus": 2, "publication_state": "baseline_only"})
    rows = data_health.run_all("authoring")
    assert len(rows) == 4
    assert rows[1].status == "warn"
    assert not any("family_catalogue.sqlite3" in row.detail for row in rows)


def test_authoring_retrieval_uses_explicit_id_not_filename(tmp_path, monkeypatch):
    from scripts import build_retrieval_cards as cards
    folder = tmp_path / "knowledge" / "test"
    (folder / "research").mkdir(parents=True)
    (folder / "families.json").write_text(json.dumps({"families": [{"family_id": "TEST", "name": "Changed name"}]}))
    (folder / "research" / "unrelated_filename.json").write_text(json.dumps({
        "family_id": "TEST", "spec": {"description": "Retained supplied description"}}))
    monkeypatch.setattr(cards, "ROOT", tmp_path)
    rows = cards.build_all("authoring")
    assert rows[0]["description"] == "Retained supplied description"
    assert rows[0]["purpose"] == "unreviewed_authoring_retrieval_not_primary_evidence"


def test_guide_loader_preserves_legacy_non_ok_record_by_identity(tmp_path, monkeypatch):
    from scripts import enrich_knowledge_docs as guides
    folder = tmp_path / "knowledge" / "test" / "research"
    folder.mkdir(parents=True)
    value = {"family_id": "TEST", "status": "original_unavailable", "spec": {"range": [{"length": "22.25 m"}]}}
    (folder / "unknown-name.json").write_text(json.dumps(value))
    monkeypatch.setattr(guides, "ROOT", tmp_path)
    assert guides.load_research("test", "Renamed family", "TEST")["spec"] == value["spec"]
