"""Lossless, owner-mapped commercial catalogue versions; new rows start held."""
from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import Path

from research_store import Conflict, canonical


def preview(root: Path, source: Path, mapping: dict, sheet: str | None = None) -> dict:
    root, source = root.resolve(), source.resolve()
    if not source.is_relative_to(root / "data") or not source.is_file():
        raise ValueError("Commercial source must be a local file inside this checkout's data directory")
    if not isinstance(mapping, dict) or not {"family_id", "our_sku", "supplier_sku", "product_name"}.issubset(mapping):
        raise ValueError("Explicit mapping needs family_id, our_sku, supplier_sku and product_name columns")
    if any(not isinstance(key, str) or not isinstance(column, str) or not column.strip()
           for key, column in mapping.items()):
        raise ValueError("Column mappings must contain nonblank strings")
    raw = source.read_bytes()
    checksum = hashlib.sha256(raw).hexdigest()
    if source.suffix.casefold() == ".csv":
        if sheet is not None:
            raise ValueError("CSV has no worksheet")
        reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
        headers = reader.fieldnames or []
        rows = list(reader)
    elif source.suffix.casefold() == ".xlsx":
        from openpyxl import load_workbook
        if not sheet:
            raise ValueError("Choose an exact worksheet; no automatic first-sheet selection")
        book = load_workbook(source, read_only=True, data_only=False)
        try:
            if sheet not in book.sheetnames:
                raise ValueError("Unknown worksheet")
            values = book[sheet].values
            headers = list(next(values, ()))
            rows = [dict(zip(headers, row)) for row in values]
        finally:
            book.close()
    else:
        raise ValueError("Only CSV and XLSX commercial input supported")
    if len(headers) != len(set(headers)) or any(not isinstance(h, str) or not h.strip() for h in headers):
        raise ValueError("Column headers must be nonblank and unique")
    if not rows or any(column not in headers for column in mapping.values()):
        raise ValueError("Mapped columns missing or input has no rows")
    if any(None in row for row in rows):
        raise ValueError("CSV contains extra values without column headers")
    families = {}
    for path in root.glob("knowledge/*/families.json"):
        families.update({r["family_id"]:r for r in json.loads(path.read_text(encoding="utf-8-sig"))["families"]})
    records, blockers = [], []
    for number, raw_row in enumerate(rows, 2):
        values = {key: "" if raw_row.get(column) is None else str(raw_row[column]) for key, column in mapping.items()}
        family = families.get(values["family_id"])
        if family is None:
            blockers.append(f"Row {number}: unknown canonical family_id {values['family_id']!r}")
        records.append({**values, "sku_record_id":f"V-{checksum[:16]}-{number}",
                        "family_name":family["name"] if family else "", "manufacturer":family.get("manufacturer","") if family else "",
                        "validation_status":"REVIEW","validation_notes":"New commercial version; exact variant continuity and evidence applicability need human review.",
                        "active":values.get("active","Unknown"),"bot_content_status":"HOLD",
                        "recommendation_eligible":"False","source_workbook":source.name,"source_sha256":checksum,
                        "source_sheet":sheet or "", "source_row":number,
                        "raw_values":{key:value if value is None or isinstance(value,(str,int,float,bool))
                                      else {"type":type(value).__name__,"value":str(value)} for key,value in raw_row.items()}})
    codes = {}
    for row in records:
        for field in ("our_sku","supplier_sku"):
            codes.setdefault((field,row[field]),[]).append(row["sku_record_id"])
    library=CatalogueLibrary(root)
    previous=library.active_rows()
    if previous is None:
        path=root/"data"/"processed"/"product_catalogue_skus.csv"
        previous=[]
        if path.is_file():
            with path.open(encoding="utf-8-sig",newline="") as handle:
                previous=list(csv.DictReader(handle))
    old_codes={(field,row.get(field,"")) for row in previous for field in ("our_sku","supplier_sku") if row.get(field)}
    new_codes={(field,row[field]) for row in records for field in ("our_sku","supplier_sku") if row[field]}
    body = {"source_sha256":checksum,"source_file":source.name,"source_path":source.relative_to(root).as_posix(),"sheet":sheet,"mapping":mapping,
            "rows":records,"blockers":blockers,
            "duplicate_codes":[{"field":field,"code":code,"row_ids":ids} for (field,code),ids in codes.items() if code and len(ids)>1],
            "blank_code_rows":[r["sku_record_id"] for r in records if not r["our_sku"] or not r["supplier_sku"]],
            "approval_transfer":False,"variant_continuity":"unreviewed",
            "changes":{"previous_version":library.active_id() or "baseline_csv","old_row_count":len(previous),
                       "new_row_count":len(records),"added_codes":[list(item) for item in sorted(new_codes-old_codes)],
                       "removed_codes":[list(item) for item in sorted(old_codes-new_codes)],
                       "unreviewed_continuity_rows":[row["sku_record_id"] for row in records]}}
    return {"version_id":hashlib.sha256(canonical(body).encode()).hexdigest(),"payload":body}


class CatalogueLibrary:
    def __init__(self, root: Path):
        self.root = root.resolve()
        self.directory = self.root / "data" / "local" / "catalogue_versions"
        if not self.directory.resolve().is_relative_to(self.root):
            raise ValueError("Version directory escapes checkout")

    def stage(self, data: dict):
        if not isinstance(data, dict) or not isinstance(data.get("payload"), dict):
            raise ValueError("Unmodified catalogue preview required")
        if not isinstance(data["payload"].get("source_path"), str):
            raise ValueError("Local catalogue source path required")
        if hashlib.sha256(canonical(data["payload"]).encode()).hexdigest() != data["version_id"]:
            raise ValueError("Catalogue version checksum mismatch")
        source=(self.root/data["payload"]["source_path"]).resolve()
        if not source.is_relative_to(self.root/"data"):
            raise ValueError("Catalogue source escapes local data library")
        original=source.read_bytes()
        if hashlib.sha256(original).hexdigest()!=data["payload"]["source_sha256"]:
            raise ValueError("Source changed after preview")
        regenerated = preview(self.root, source, data["payload"]["mapping"], data["payload"].get("sheet"))
        if canonical(data) != canonical(regenerated):
            raise ValueError("Catalogue preview changed or is stale; preview the source again")
        self.directory.mkdir(parents=True, exist_ok=True)
        archive=self.directory/(data["version_id"][:16]+".source"+source.suffix.casefold())
        if archive.exists():
            if archive.read_bytes()!=original:
                raise ValueError("Original-source archive collision")
        else:
            with archive.open("xb") as handle:
                handle.write(original)
        target = self.directory / (data["version_id"][:16]+".json")
        with target.open("x",encoding="utf-8") as handle:
            handle.write(canonical(data))
        return target

    def staged(self):
        """Read-only listing of every recorded catalogue version (staged or active). No mutation, no approval."""
        if not self.directory.is_dir():
            return []
        active = self.active_id()
        items = []
        for path in sorted(self.directory.glob("*.json")):
            if path.name == "active.json":
                continue
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
            except (json.JSONDecodeError, OSError) as exc:
                raise ValueError(f"Cannot read catalogue version {path.name}: {exc}") from exc
            if not isinstance(data, dict):
                raise ValueError(f"Invalid catalogue version record: {path.name}")
            version_id = data.get("version_id")
            payload = data.get("payload")
            if not isinstance(version_id, str) or not isinstance(payload, dict):
                raise ValueError(f"Invalid catalogue version record: {path.name}")
            try:
                data = self.read(version_id)
            except (OSError, ValueError, KeyError) as exc:
                raise ValueError(f"Invalid catalogue version {path.name}: {exc}") from exc
            if self.version_path(version_id) != path:
                raise ValueError(f"Catalogue version filename mismatch: {path.name}")
            changes = payload.get("changes", {})
            items.append({
                "version_id": version_id,
                "active": version_id == active,
                "source_file": payload.get("source_file"),
                "source_path": payload.get("source_path"),
                "source_sha256": payload.get("source_sha256"),
                "row_count": len(payload.get("rows", [])),
                "blockers": payload.get("blockers", []),
                "duplicate_codes": payload.get("duplicate_codes", []),
                "blank_code_rows": payload.get("blank_code_rows", []),
                "variant_continuity": payload.get("variant_continuity"),
                "previous_version": changes.get("previous_version"),
                "added_codes": changes.get("added_codes", []),
                "removed_codes": changes.get("removed_codes", []),
                "unreviewed_continuity_rows": changes.get("unreviewed_continuity_rows", []),
            })
        return items

    def active_id(self):
        pointer=self.directory/"active.json"
        return json.loads(pointer.read_text())["version_id"] if pointer.exists() else None

    def active_rows(self):
        identifier=self.active_id()
        if identifier is None:
            return None
        return self.read(identifier)["payload"]["rows"]

    def read(self, identifier):
        from knowledge_release import RELEASE_ID
        if not isinstance(identifier,str) or not RELEASE_ID.fullmatch(identifier):
            raise ValueError("Invalid catalogue version ID")
        data=json.loads(self.version_path(identifier).read_text(encoding="utf-8"))
        if data["version_id"] != identifier or hashlib.sha256(canonical(data["payload"]).encode()).hexdigest() != identifier:
            raise ValueError("Catalogue version checksum mismatch")
        return data

    def version_path(self, identifier):
        from knowledge_release import RELEASE_ID
        if not isinstance(identifier,str) or not RELEASE_ID.fullmatch(identifier):
            raise ValueError("Invalid catalogue version ID")
        short=self.directory/(identifier[:16]+".json")
        return short if short.exists() else self.directory/(identifier+".json")

    def activate(self, identifier, confirm, expected):
        import os
        import uuid
        data=self.read(identifier)
        if confirm != identifier or data["payload"]["blockers"]:
            raise ValueError("Exact approval and resolved family mappings required")
        lock=self.directory/"activation.lock"
        with lock.open("x") as handle:
            handle.write(str(os.getpid()))
        temporary=self.directory/("a-"+uuid.uuid4().hex[:12]+".tmp")
        try:
            if self.active_id()!=expected:
                raise Conflict("Active catalogue changed since preview")
            with temporary.open("x",encoding="utf-8") as handle:
                handle.write(canonical({"version_id":identifier}))
            os.replace(temporary,self.directory/"active.json")
        finally:
            temporary.unlink(missing_ok=True)
            lock.unlink()
